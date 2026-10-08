"""Spreadsheet/CSV summaries: sheets, headers, row counts, sample rows, formulas."""

from __future__ import annotations

from pathlib import Path

import pandas as pd

SAMPLE_ROWS = 15


def _frame_md(df: pd.DataFrame, rows: int) -> str:
    sample = df.head(rows)
    return sample.to_markdown(index=False) if not sample.empty else "_(no data rows)_"


def _describe(df: pd.DataFrame) -> str:
    parts = []
    for col in df.columns:
        s = df[col]
        kind = (
            "number"
            if pd.api.types.is_numeric_dtype(s)
            else ("date" if pd.api.types.is_datetime64_any_dtype(s) else "text")
        )
        extra = ""
        if kind == "number" and s.notna().any():
            extra = f", min {s.min():g}, max {s.max():g}, sum {s.sum():g}"
        parts.append(f"- `{col}`: {kind}, {int(s.notna().sum())} values{extra}")
    return "\n".join(parts)


def read_csv(path: Path, sep: str | None = None, sample_rows: int = SAMPLE_ROWS) -> str:
    df = pd.read_csv(path, sep=sep, engine="python", encoding_errors="replace")
    return (
        f"## {path.name}\n\nRows: {len(df)}, columns: {len(df.columns)}\n\n"
        f"### Columns\n{_describe(df)}\n\n### First {min(sample_rows, len(df))} rows\n{_frame_md(df, sample_rows)}\n"
    )


def read_excel(path: Path, sheet: str | None = None, sample_rows: int = SAMPLE_ROWS) -> str:
    out = []
    formula_counts: dict[str, int] = {}
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=False)
        for ws in wb.worksheets:
            n = 0
            for row in ws.iter_rows(values_only=True):
                n += sum(1 for v in row if isinstance(v, str) and v.startswith("="))
            formula_counts[ws.title] = n
        wb.close()
    sheets = pd.read_excel(path, sheet_name=sheet if sheet else None)
    if isinstance(sheets, pd.DataFrame):
        sheets = {sheet or "Sheet1": sheets}
    out.append(f"# {path.name}\n\nSheets: {', '.join(str(s) for s in sheets)}\n")
    for name, df in sheets.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        fc = formula_counts.get(str(name))
        out.append(
            f"## Sheet: {name}\n\nRows: {len(df)}, columns: {len(df.columns)}"
            + (f", formulas: {fc}" if fc else "")
            + "\n"
        )
        out.append(f"### Columns\n{_describe(df)}\n")
        out.append(f"### First {min(sample_rows, len(df))} rows\n{_frame_md(df, sample_rows)}\n")
    return "\n".join(out)
