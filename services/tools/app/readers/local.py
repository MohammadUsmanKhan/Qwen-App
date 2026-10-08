"""Fallback extractors used when Docling is unavailable (no OCR, simpler layout)."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    out: list[str] = []
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph

            p = Paragraph(block, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading") and style[-1:].isdigit():
                out.append("#" * int(style[-1]) + " " + text)
            elif style == "Title":
                out.append("# " + text)
            elif "List" in style:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            from docx.table import Table

            t = Table(block, doc)
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in t.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * len(rows[0]))
                out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
        out.append("")
    return "\n".join(out)


def read_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []

    def shape_text(shape: Any) -> list[str]:
        lines: list[str] = []
        if shape.shape_type == 6:  # group
            for s in shape.shapes:
                lines.extend(shape_text(s))
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            lines.append(shape.text_frame.text.strip())
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        return lines

    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            out.extend(shape_text(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"> Notes: {slide.notes_slide.notes_text_frame.text.strip()}")
        out.append("")
    return "\n".join(out)


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append(f"## Page {i}\n\n{(page.extract_text() or '').strip()}")
    return "\n\n".join(pages)


def read_html(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_bytes(), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "noscript"]):
        tag.decompose()
    return soup.get_text("\n", strip=True)


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")
