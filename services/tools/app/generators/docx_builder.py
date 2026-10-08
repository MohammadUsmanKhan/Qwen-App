"""Render a DocumentSpec to .docx with python-docx.

Layout comes entirely from code and styles; the model only supplies content.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .doc_spec import (
    Block,
    CalloutBlock,
    DocStyle,
    DocumentSpec,
    HeadingBlock,
    ImageBlock,
    ListBlock,
    ListItem,
    MarkdownBlock,
    PageBreakBlock,
    ParagraphBlock,
    QuoteBlock,
    TableBlock,
)
from .markdown_blocks import markdown_to_blocks

ImageResolver = Callable[[str], Path]

PAGE_SIZES_CM = {"A4": (21.0, 29.7), "Letter": (21.59, 27.94)}
MARGIN_CM = 2.2
ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
TONES = {"info": "2E75B6", "note": "7F7F7F", "warning": "C55A11", "success": "548235"}
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\))")
_NUMBER = re.compile(r"^-?[\d,]*\.?\d+%?$")


@dataclass
class BuildResult:
    data: bytes
    headings: list[tuple[int, str]] = field(default_factory=list)
    tables: int = 0
    images: int = 0
    words: int = 0
    warnings: list[str] = field(default_factory=list)


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.upper())


def _tint(hex_: str, amount: float) -> str:
    """Mix a colour with white; amount=0.85 → very light."""
    r, g, b = (int(hex_[i : i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c + (255 - c) * amount):02X}" for c in (r, g, b))


def _set_cell_shading(cell: _Cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_borders(cell: _Cell, **edges: tuple[str, int]) -> None:
    """edges: left=("1F4E79", 24) → colour, width in eighths of a point."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            color, size = edges[edge]
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)


def _add_field(paragraph: Paragraph, instr: str, placeholder: str = "1") -> None:
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = placeholder
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


class DocxBuilder:
    def __init__(self, spec: DocumentSpec, resolve_image: ImageResolver | None = None) -> None:
        self.spec = spec
        self.style: DocStyle = spec.style
        self.resolve_image = resolve_image
        self.doc: DocxDocument = Document()
        self.result = BuildResult(data=b"")

    # ------------------------------------------------------------------ public
    def build(self) -> BuildResult:
        blocks = self._expand(self.spec.blocks)
        self.result.headings = [(b.level, b.text) for b in blocks if isinstance(b, HeadingBlock)]
        self._setup_page()
        self._setup_styles()
        self._properties()
        self._title_area()
        if self.spec.toc:
            self._toc(self.result.headings)
        for block in blocks:
            self._block(block)
        self._header_footer()
        buf = io.BytesIO()
        self.doc.save(buf)
        self.result.data = buf.getvalue()
        return self.result

    # ------------------------------------------------------------------ setup
    @staticmethod
    def _expand(blocks: list[Block]) -> list[Block]:
        out: list[Block] = []
        for b in blocks:
            out.extend(markdown_to_blocks(b.text) if isinstance(b, MarkdownBlock) else [b])
        return out

    def _setup_page(self) -> None:
        w, h = PAGE_SIZES_CM[self.spec.page_size]
        sec = self.doc.sections[0]
        if self.spec.orientation == "landscape":
            w, h = h, w
            sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(w), Cm(h)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(MARGIN_CM))

    @property
    def text_width_cm(self) -> float:
        sec = self.doc.sections[0]
        width = int(sec.page_width or 0) - int(sec.left_margin or 0) - int(sec.right_margin or 0)
        return width / 360000

    def _font(
        self,
        style_name: str,
        *,
        font: str,
        size: float | None = None,
        color: str | None = None,
        bold: bool | None = None,
    ) -> None:
        st = self.doc.styles[style_name]
        st.font.name = font
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        if size:
            st.font.size = Pt(size)
        if color:
            st.font.color.rgb = _rgb(color)
        if bold is not None:
            st.font.bold = bold

    def _setup_styles(self) -> None:
        s = self.style
        body = s.body_size_pt
        self._font("Normal", font=s.body_font, size=body, color=s.text_color)
        normal = self.doc.styles["Normal"].paragraph_format
        normal.space_after = Pt(6)
        normal.line_spacing = 1.15
        self._font("Title", font=s.heading_font, size=body * 2.6, color=s.accent_color, bold=True)
        self._font("Subtitle", font=s.heading_font, size=body * 1.4, color="595959")
        for level, scale in ((1, 1.65), (2, 1.3), (3, 1.12)):
            name = f"Heading {level}"
            self._font(name, font=s.heading_font, size=body * scale, color=s.accent_color, bold=True)
            pf = self.doc.styles[name].paragraph_format
            pf.space_before = Pt(18 if level == 1 else 12)
            pf.space_after = Pt(6)
            pf.keep_with_next = True
        self._font("Caption", font=s.body_font, size=body - 1.5, color="595959", bold=False)
        self.doc.styles["Caption"].font.italic = True
        self._font("Quote", font=s.body_font, size=body, color="404040")
        for name in ("Header", "Footer"):
            self._font(name, font=s.body_font, size=9, color="7F7F7F")
        for level in (1, 2, 3):
            name = f"TOC {level}"
            if name not in [st.name for st in self.doc.styles]:
                st = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = self.doc.styles["Normal"]
            pf = self.doc.styles[name].paragraph_format
            pf.left_indent = Cm(0.6 * (level - 1))
            pf.space_after = Pt(2)
            if level == 1:
                self.doc.styles[name].font.bold = True

    def _properties(self) -> None:
        cp = self.doc.core_properties
        cp.title = self.spec.title
        cp.subject = self.spec.subtitle or ""
        cp.author = self.spec.author or ""
        cp.last_modified_by = self.spec.author or ""
        cp.comments = ""

    # ------------------------------------------------------------------ title / toc
    def _title_area(self) -> None:
        spec = self.spec
        if spec.cover_page:
            self.doc.sections[0].different_first_page_header_footer = True
            for _ in range(6):
                self.doc.add_paragraph()
            if spec.logo_file_id:
                self._picture(spec.logo_file_id, width_cm=5, align="left")
        self.doc.add_paragraph(spec.title, style="Title")
        if spec.subtitle:
            self.doc.add_paragraph(spec.subtitle, style="Subtitle")
        meta = " · ".join(x for x in (spec.author, spec.date) if x)
        if meta:
            p = self.doc.add_paragraph()
            r = p.add_run(meta)
            r.font.color.rgb = _rgb("595959")
        if spec.cover_page:
            self._page_break()

    def _toc(self, headings: list[tuple[int, str]]) -> None:
        self.doc.add_paragraph("Contents", style="Heading 1")
        entries = headings or [(1, "Right-click and choose Update Field to build the table of contents")]
        # One TOC field spanning several paragraphs; the cached entries show the
        # headings before Word refreshes it (page numbers appear after the refresh).
        first, last = 0, len(entries) - 1
        for i, (level, text) in enumerate(entries):
            p = self.doc.add_paragraph(style=f"TOC {level}")
            if i == first:
                self._fld_char(p, "begin")
                instr = OxmlElement("w:r")
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-3" \\h \\z \\u '
                instr.append(it)
                p._p.append(instr)
                self._fld_char(p, "separate")
            p.add_run(text)
            if i == last:
                self._fld_char(p, "end")
        settings = self.doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)
        self._page_break()

    @staticmethod
    def _fld_char(p: Paragraph, kind: str) -> None:
        r = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        r.append(fc)
        p._p.append(r)

    def _header_footer(self) -> None:
        sec = self.doc.sections[0]
        if self.spec.header_text:
            p = sec.header.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run(self.spec.header_text)
        if self.spec.footer_text or self.spec.page_numbers:
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if self.spec.footer_text:
                p.add_run(self.spec.footer_text + ("   ·   " if self.spec.page_numbers else ""))
            if self.spec.page_numbers:
                p.add_run("Page ")
                _add_field(p, "PAGE")
                p.add_run(" of ")
                _add_field(p, "NUMPAGES")

    # ------------------------------------------------------------------ blocks
    def _block(self, b: Block) -> None:
        if isinstance(b, HeadingBlock):
            self.doc.add_paragraph(b.text, style=f"Heading {b.level}")
        elif isinstance(b, ParagraphBlock):
            p = self.doc.add_paragraph()
            p.alignment = ALIGN[b.align]
            self._inline(p, b.text)
        elif isinstance(b, ListBlock):
            self._list(b)
        elif isinstance(b, TableBlock):
            self._table(b)
        elif isinstance(b, ImageBlock):
            self._picture(b.file_id, b.width_cm, caption=b.caption)
        elif isinstance(b, QuoteBlock):
            p = self.doc.add_paragraph(style="Quote")
            self._inline(p, b.text)
            if b.attribution:
                a = self.doc.add_paragraph(f"— {b.attribution}", style="Quote")
                a.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif isinstance(b, CalloutBlock):
            self._callout(b)
        elif isinstance(b, PageBreakBlock):
            self._page_break()

    def _page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _inline(self, p: Paragraph, text: str) -> None:
        self.result.words += len(text.split())
        for part in _INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`") and len(part) > 2:
                r = p.add_run(part[1:-1])
                r.font.name = "Consolas"
            elif part.startswith("[") and "](" in part and part.endswith(")"):
                label, url = part[1:-1].split("](", 1)
                self._hyperlink(p, label, url)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                p.add_run(part[1:-1]).italic = True
            else:
                p.add_run(part)

    def _hyperlink(self, p: Paragraph, text: str, url: str) -> None:
        r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), self.style.accent_color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.extend([color, underline])
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        run.extend([rpr, t])
        link.append(run)
        p._p.append(link)

    def _new_numbering(self) -> int:
        """A fresh numbering instance so each numbered list restarts at 1."""
        numbering = self.doc.part.numbering_part.element
        style_num_id = self.doc.styles["List Number"].element.pPr.numPr.numId.val
        abstract_id = numbering.num_having_numId(style_num_id).abstractNumId.val
        num = numbering.add_num(abstract_id)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return int(num.numId)

    def _list(self, b: ListBlock) -> None:
        num_id = self._new_numbering() if b.style == "number" else None
        for raw in b.items:
            item = raw if isinstance(raw, ListItem) else ListItem(text=raw, level=0)
            if b.style == "bullet":
                style = "List Bullet" if item.level == 0 else f"List Bullet {item.level + 1}"
                p = self.doc.add_paragraph(style=style)
            else:
                p = self.doc.add_paragraph(style="List Number")
                num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_ilvl().val = item.level
                num_pr.get_or_add_numId().val = num_id
                if item.level:
                    p.paragraph_format.left_indent = Cm(0.63 * (item.level + 1))
            p.paragraph_format.space_after = Pt(3)
            self._inline(p, item.text)

    def _table(self, b: TableBlock) -> None:
        self.result.tables += 1
        if b.caption:
            self.doc.add_paragraph(b.caption, style="Caption")
        table: Table = self.doc.add_table(rows=1 + len(b.rows), cols=len(b.columns))
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        accent, band = self.style.accent_color, _tint(self.style.accent_color, 0.9)
        numeric = [
            all(_is_number(r[c]) for r in b.rows if r[c] not in (None, ""))
            and any(r[c] not in (None, "") for r in b.rows)
            for c in range(len(b.columns))
        ]

        header = table.rows[0]
        tr_pr = header._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)
        for c, name in enumerate(b.columns):
            cell = header.cells[c]
            _set_cell_shading(cell, accent)
            p = cell.paragraphs[0]
            run = p.add_run(name)
            run.bold, run.font.color.rgb = True, _rgb("FFFFFF")
            if numeric[c]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r, row in enumerate(b.rows, start=1):
            for c, value in enumerate(row):
                cell = table.rows[r].cells[c]
                if r % 2 == 0:
                    _set_cell_shading(cell, band)
                p = cell.paragraphs[0]
                self._inline(p, _fmt(value))
                if numeric[c]:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
        _set_table_borders(table, "BFBFBF")
        widths = b.column_widths_cm or _auto_widths(b, self.text_width_cm)
        total = sum(widths)
        scale = min(1.0, self.text_width_cm / total) if total else 1.0
        table.autofit = False
        for c, w in enumerate(widths):
            table.columns[c].width = Cm(w * scale)
            for row in table.rows:
                row.cells[c].width = Cm(w * scale)
        self.doc.add_paragraph()

    def _picture(self, file_id: str, width_cm: float, caption: str | None = None, align: str = "center") -> None:
        if self.resolve_image is None:
            self.result.warnings.append(f"image {file_id}: no image source configured")
            return
        try:
            path = self.resolve_image(file_id)
            self.doc.add_picture(str(path), width=Cm(min(width_cm, self.text_width_cm)))
        except Exception as e:  # noqa: BLE001 — report and keep building
            self.result.warnings.append(f"image {file_id} skipped: {e}")
            return
        self.result.images += 1
        self.doc.paragraphs[-1].alignment = ALIGN[align]
        if caption:
            cap = self.doc.add_paragraph(caption, style="Caption")
            cap.alignment = ALIGN[align]

    def _callout(self, b: CalloutBlock) -> None:
        color = TONES[b.tone]
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        _set_cell_shading(cell, _tint(color, 0.88))
        _set_cell_borders(cell, left=(color, 24))
        p = cell.paragraphs[0]
        if b.title:
            r = p.add_run(b.title)
            r.bold, r.font.color.rgb = True, _rgb(color)
            p = cell.add_paragraph()
        self._inline(p, b.text)
        self.doc.add_paragraph()


def _set_table_borders(table: Table, color: str) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def _auto_widths(b: TableBlock, available_cm: float) -> list[float]:
    """Share the page width by content length, so short numeric columns stay narrow."""
    lengths = []
    for c, name in enumerate(b.columns):
        longest = max([len(name)] + [len(_fmt(r[c])) for r in b.rows])
        lengths.append(min(max(longest, 4), 40))
    total = sum(lengths)
    return [available_cm * n / total for n in lengths]


def _is_number(v: object) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, int | float):
        return True
    s = str(v).strip().replace(" ", "")
    for sym in ("$", "€", "£", "¥"):
        s = s.replace(sym, "")
    return bool(_NUMBER.match(s))


def _fmt(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.2f}".rstrip("0").rstrip(".")
    return str(v)


def build_docx(spec: DocumentSpec, resolve_image: ImageResolver | None = None) -> BuildResult:
    return DocxBuilder(spec, resolve_image).build()
