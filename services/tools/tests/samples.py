"""Build small sample files in memory for reader tests."""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Quarterly Review", 1)
    doc.add_paragraph("Revenue grew in every region.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Region", "Revenue"
    t.cell(1, 0).text, t.cell(1, 1).text = "North", "120"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Company A Services"
    slide.placeholders[1].text = "Consulting\nTraining"
    slide.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(0.5)).text_frame.text = "info@company-a.com"
    slide.notes_slide.notes_text_frame.text = "Speaker note here"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Qty", "Price", "Total"])
    for i, (item, q, p) in enumerate([("Widget", 3, 2.5), ("Gadget", 1, 10.0)], start=2):
        ws.append([item, q, p, f"=B{i}*C{i}"])
    wb.create_sheet("Notes").append(["Prepared by finance"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
