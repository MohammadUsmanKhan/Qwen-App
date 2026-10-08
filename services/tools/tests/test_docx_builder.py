from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pydantic import ValidationError

from app.generators.doc_spec import DocumentSpec
from app.generators.docx_builder import build_docx


def spec(**kw: object) -> DocumentSpec:
    base: dict[str, object] = {"title": "Site Survey", "blocks": []}
    base.update(kw)
    return DocumentSpec.model_validate(base)


def open_doc(data: bytes):  # type: ignore[no-untyped-def]
    return Document(io.BytesIO(data))


def test_headings_paragraphs_and_properties() -> None:
    res = build_docx(
        spec(
            author="Acme",
            subtitle="Riverside",
            blocks=[
                {"type": "heading", "text": "Scope", "level": 1},
                {"type": "paragraph", "text": "Covers **three** buildings and [our site](https://example.com)."},
                {"type": "heading", "text": "Detail", "level": 2},
            ],
        )
    )
    doc = open_doc(res.data)
    styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Title", "Site Survey") in styles
    assert ("Heading 1", "Scope") in styles and ("Heading 2", "Detail") in styles
    para = next(p for p in doc.paragraphs if p.text.startswith("Covers"))
    assert any(r.bold and r.text == "three" for r in para.runs)
    assert any("example.com" in r.target_ref for r in doc.part.rels.values() if r.is_external)
    assert doc.core_properties.title == "Site Survey" and doc.core_properties.author == "Acme"
    assert res.headings == [(1, "Scope"), (2, "Detail")]


def test_numbered_lists_restart() -> None:
    res = build_docx(
        spec(
            blocks=[
                {"type": "list", "style": "number", "items": ["a", "b"]},
                {"type": "paragraph", "text": "between"},
                {"type": "list", "style": "number", "items": ["c", {"text": "c.1", "level": 1}]},
            ]
        )
    )
    doc = open_doc(res.data)
    num_ids = []
    for p in doc.paragraphs:
        num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
        if num_pr is not None and num_pr.numId is not None:
            num_ids.append((p.text, num_pr.numId.val, num_pr.ilvl.val))
    assert len({num for text, num, _lvl in num_ids if text in ("a", "b")}) == 1
    first, second = num_ids[0][1], num_ids[2][1]
    assert first != second, "second numbered list must restart with its own numbering"
    assert ("c.1", second, 1) in num_ids


def test_table_header_numbers_and_padding() -> None:
    res = build_docx(
        spec(
            blocks=[
                {
                    "type": "table",
                    "caption": "Costs",
                    "columns": ["Item", "Qty", "Cost"],
                    "rows": [["Roof", 1, 12500.5], ["Windows", 24]],
                }
            ]
        )
    )
    doc = open_doc(res.data)
    t = doc.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["Item", "Qty", "Cost"]
    assert [c.text for c in t.rows[1].cells] == ["Roof", "1", "12500.5"]
    assert t.rows[2].cells[2].text == ""  # short row padded
    assert t.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None  # header repeats on each page
    assert res.tables == 1


def test_table_row_too_long_is_rejected() -> None:
    with pytest.raises(ValidationError, match="row 0 has 3 values"):
        spec(blocks=[{"type": "table", "columns": ["A", "B"], "rows": [[1, 2, 3]]}])


def test_cover_toc_header_footer() -> None:
    res = build_docx(
        spec(
            cover_page=True,
            toc=True,
            header_text="Confidential",
            footer_text="Acme",
            blocks=[{"type": "heading", "text": "Intro", "level": 1}],
        )
    )
    doc = open_doc(res.data)
    xml = doc.element.xml
    assert 'TOC \\o "1-3"' in xml
    toc_entries = [p.text for p in doc.paragraphs if p.style.name.startswith("TOC")]
    assert toc_entries == ["Intro"]
    sec = doc.sections[0]
    assert sec.different_first_page_header_footer
    assert sec.header.paragraphs[0].text == "Confidential"
    footer_xml = sec.footer._element.xml
    assert "Acme" in footer_xml and "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_markdown_block_expands() -> None:
    md = "# Plan\n\nIntro text\nwrapped.\n\n- one\n  - nested\n1. first\n2. second\n\n| A | B |\n|---|---|\n| 1 | 2 |"
    res = build_docx(spec(blocks=[{"type": "markdown", "text": md}]))
    doc = open_doc(res.data)
    texts = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Heading 1", "Plan") in texts
    assert ("Normal", "Intro text wrapped.") in texts
    assert ("List Bullet 2", "nested") in texts
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["1", "2"]


def test_images_resolved_and_missing_reported(tmp_path: Path) -> None:
    img = tmp_path / "chart.png"
    Image.new("RGB", (200, 100), "red").save(img)

    def resolve(file_id: str) -> Path:
        if file_id == "gen_ok":
            return img
        raise FileNotFoundError(file_id)

    res = build_docx(
        spec(
            blocks=[
                {"type": "image", "file_id": "gen_ok", "caption": "Figure 1"},
                {"type": "image", "file_id": "gen_missing"},
            ]
        ),
        resolve,
    )
    assert res.images == 1
    assert len(res.warnings) == 1 and "gen_missing" in res.warnings[0]
    doc = open_doc(res.data)
    assert len(doc.inline_shapes) == 1
    assert any(p.text == "Figure 1" and p.style.name == "Caption" for p in doc.paragraphs)


def test_style_colours_validated() -> None:
    with pytest.raises(ValidationError):
        spec(style={"accent_color": "blue"})
    res = build_docx(
        spec(
            style={"accent_color": "#c00000", "heading_font": "Georgia"},
            blocks=[{"type": "heading", "text": "H", "level": 1}],
        )
    )
    h1 = open_doc(res.data).styles["Heading 1"]
    assert str(h1.font.color.rgb) == "C00000" and h1.font.name == "Georgia"


def test_landscape_letter() -> None:
    doc = open_doc(build_docx(spec(page_size="Letter", orientation="landscape")).data)
    sec = doc.sections[0]
    assert sec.page_width > sec.page_height
